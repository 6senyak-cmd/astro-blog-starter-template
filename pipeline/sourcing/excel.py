"""
Excel output.

Two sheets. The pipeline, and the methodology behind it.

The methodology sheet is not decoration. Every ARR figure in this workbook
is an estimate, and estimates lose their caveats the moment they are
copied into a deck. Keeping the basis in a column next to the number --
and the benchmarks on a sheet anyone can open -- is what stops "proxy,
low confidence" from becoming "$4.2M ARR" in an IC memo three weeks later.
"""

from __future__ import annotations

from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
INPUT_FONT = Font(name=FONT, size=10, color="0000FF")

CONF_FILL = {
    "high": PatternFill("solid", fgColor="C6EFCE"),
    "medium": PatternFill("solid", fgColor="FFEB9C"),
    "low": PatternFill("solid", fgColor="FFC7CE"),
    "none": PatternFill("solid", fgColor="F2F2F2"),
}

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    ("Company", 34),
    ("Fit", 6),
    ("ARR estimate", 18),
    ("Confidence", 12),
    ("Method", 11),
    ("ARR basis", 52),
    ("Raised ($)", 14),
    ("Founded", 9),
    ("Founder-led", 12),
    ("Founder evidence", 34),
    ("Model", 20),
    ("Open roles", 11),
    ("Commercial roles", 16),
    ("Growth", 24),
    ("Rationale", 62),
    ("Disqualifiers", 30),
    ("Location", 18),
    ("Officers (Form D)", 42),
    ("Form D", 10),
    ("Job board", 11),
    ("First seen", 12),
]


def _link(cell, url: str, label: str) -> None:
    if not url:
        return
    cell.value = label
    cell.hyperlink = url
    cell.font = Font(name=FONT, size=10, color="0563C1", underline="single")


def write(rows: list[dict], cfg: dict, path: str) -> str:
    rows = sorted(rows, key=lambda r: r.get("fit_score") or 0, reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Pipeline"

    for idx, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows)+1, 2)}"

    for r, row in enumerate(rows, start=2):
        values = [
            row.get("entity_name"),
            row.get("fit_score"),
            row.get("arr_band"),
            row.get("arr_confidence"),
            row.get("arr_method"),
            row.get("arr_basis"),
            row.get("total_sold_usd"),
            row.get("year_incorporated"),
            {True: "Yes", False: "No", None: "Unknown"}.get(row.get("founder_led"), "Unknown"),
            row.get("founder_evidence"),
            row.get("business_model"),
            row.get("open_roles"),
            row.get("commercial_roles"),
            row.get("growth_flag"),
            row.get("rationale"),
            "; ".join(row.get("disqualifiers") or []),
            f"{row.get('city','')}, {row.get('state','')}".strip(", "),
            row.get("officers"),
            None,
            None,
            row.get("first_seen"),
        ]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=c in (6, 10, 15, 16, 18))

        ws.cell(row=r, column=4).fill = CONF_FILL.get(row.get("arr_confidence"), CONF_FILL["none"])
        ws.cell(row=r, column=7).number_format = '$#,##0;($#,##0);-'
        ws.cell(row=r, column=8).number_format = "0"
        _link(ws.cell(row=r, column=19), row.get("filing_url", ""), "filing")
        _link(ws.cell(row=r, column=20), row.get("board_url", ""), "careers")
        ws.row_dimensions[r].height = 42

    _methodology(wb, cfg, len(rows))
    wb.save(path)
    return path


def _methodology(wb: Workbook, cfg: dict, n_rows: int) -> None:
    ws = wb.create_sheet("Methodology")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 78

    model = cfg["arr_model"]
    blocks = [
        ("Generated", date.today().isoformat()),
        ("Companies in sheet", n_rows),
        ("Thesis", cfg["thesis"]["description"].strip()),
        ("", ""),
        ("SOURCES", ""),
        ("Discovery", "SEC EDGAR Form D filings, quarterly master index. Free, public, no API key."),
        ("Hiring signals", "Greenhouse / Lever / Ashby public job board endpoints. Unauthenticated, read-only."),
        ("Scoring", f"{cfg['scoring']['model']} — fit, founder-led judgment, and rationale "
                    "assigned interactively during the morning-source review, 50 companies per batch."),
        ("Not used", "LinkedIn, and any scraping of licensed databases. Both breach terms of service."),
        ("", ""),
        ("ARR ESTIMATION", ""),
        ("Method: disclosed", "Form D Item 5 self-reported revenue range. Highest confidence — a federal filing."),
        ("Method: stated", "Explicit revenue figure found in company job postings. Unverified marketing copy."),
        ("Method: proxy", "Headcount implied by open roles x ARR-per-FTE benchmark. Wide bands. Treat as a sort key, not a number."),
        ("", ""),
        ("ASSUMPTIONS", "Edit these in config.yaml, not here"),
        ("ARR per FTE — low", model["arr_per_fte_low"]),
        ("ARR per FTE — median", model["arr_per_fte_mid"]),
        ("ARR per FTE — high", model["arr_per_fte_high"]),
        ("Benchmark source", "SaaS Capital 2025 survey: median revenue/employee for private SaaS $129,724; "
                             "$94,444 for equity-backed companies in the $1-3M ARR band; $110,000 bootstrapped. "
                             "Metric has risen every year since 2022."),
        ("Open roles as % of headcount", f"{model['open_roles_pct_of_headcount_low']:.0%} to {model['open_roles_pct_of_headcount_high']:.0%}"),
        ("", ""),
        ("KNOWN LIMITATIONS", ""),
        ("Bootstrapped companies", "A company that never raised outside capital never files Form D and will not appear here. "
                                   "That is a real hole in the discovery layer, and it overlaps with exactly the "
                                   "capital-efficient profile in the thesis."),
        ("Growth signal on run 1", "Velocity requires two snapshots. The first run writes a baseline and reports no growth flags."),
        ("Job board coverage", "Only companies using Greenhouse, Lever, or Ashby with a public board. "
                               "Board tokens are guessed from company names, so some are missed."),
        ("Revenue non-disclosure", "Many Form D filers select 'Decline to Disclose'. Those fall back to the proxy method."),
    ]

    for r, (label, value) in enumerate(blocks, start=1):
        a = ws.cell(row=r, column=1, value=label)
        a.font = Font(name=FONT, size=10, bold=bool(label and label.isupper()))
        b = ws.cell(row=r, column=2, value=value)
        b.font = INPUT_FONT if label.startswith("ARR per FTE") else BODY_FONT
        b.alignment = Alignment(vertical="top", wrap_text=True)
