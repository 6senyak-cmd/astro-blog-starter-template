"""
Offline smoke test: exercises parse -> rules -> ARR -> workbook without
touching the network. Run from pipeline/:  python -m tests.smoke
"""

from __future__ import annotations

import sys
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from sourcing import arr, edgar, jobs, rules, excel  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
FIXTURES = Path(__file__).resolve().parent / "fixtures"


def main() -> None:
    cfg = yaml.safe_load((ROOT / "config.yaml").read_text(encoding="utf-8"))
    dcfg = cfg["discovery"]

    # A real operating SaaS company passes the rules.
    saas = edgar.parse_form_d((FIXTURES / "formd_saas.xml").read_text())
    assert saas is not None
    assert saas["entity_name"] == "Ledgerline Software Inc."
    assert saas["total_sold_usd"] == 9_500_000
    assert saas["revenue_range"] == "$1,000,001 - $5,000,000"
    assert "Dana Whitfield (Executive Officer, Director)" in saas["officers"]
    assert rules.disqualify(saas, dcfg) == []

    # A venture fund is thrown out for two independent reasons.
    fund = edgar.parse_form_d((FIXTURES / "formd_fund.xml").read_text())
    reasons = rules.disqualify(fund, dcfg)
    assert any("3C" in r for r in reasons), reasons
    assert any("pooled" in r for r in reasons), reasons

    # ARR: disclosed beats proxy.
    est = arr.estimate(saas, cfg["arr_model"])
    assert est["arr_method"] == "disclosed"
    assert est["arr_confidence"] == "high"

    # ARR: proxy fires when only hiring data exists, and the band is wide.
    proxy_row = {"open_roles": 12}
    est = arr.estimate(proxy_row, cfg["arr_model"])
    assert est["arr_method"] == "proxy"
    assert "Sort key" in est["arr_basis"]

    # ARR: stated beats proxy.
    stated_row = {"open_roles": 12, "stated_arr": {"usd": 4_000_000, "quote": "$4M in ARR"}}
    est = arr.estimate(stated_row, cfg["arr_model"])
    assert est["arr_method"] == "stated"

    # Job-board helpers: token guessing and stated-ARR extraction.
    assert jobs.token_variants("Ledgerline Software Inc.") == ["ledgerlinesoftware", "ledgerline-software", "ledgerline"]
    m = jobs.ARR_PATTERN.search("we just passed $4.5M in ARR and are growing")
    assert m and m.group(1) == "4.5"

    hit = {
        "provider": "greenhouse", "token": "ledgerline",
        "board_url": "https://boards.greenhouse.io/ledgerline", "board_verified": True,
        "postings": [
            ("Account Executive", "help us grow past $4.5M ARR"),
            ("Senior Backend Engineer", "python, postgres"),
            ("Customer Success Manager", ""),
        ],
    }
    summary = jobs.summarize(hit, cfg["jobs"]["commercial_title_keywords"])
    assert summary["open_roles"] == 3
    assert summary["commercial_roles"] == 2
    assert summary["stated_arr"]["usd"] == 4_500_000

    # Full row through the workbook writer.
    saas.update(summary)
    saas.update(arr.estimate(saas, cfg["arr_model"]))
    saas.update(
        accession="0001999999-26-000042",
        filing_url="https://www.sec.gov/Archives/edgar/data/1999999/000199999926000042/0001999999-26-000042-index.htm",
        date_filed="2026-05-20", first_seen="2026-07-19",
        fit_score=8, founder_led=True,
        founder_evidence="Two executive officers, one also a director; company incorporated 2022.",
        business_model="B2B SaaS", rationale="Disclosed $1–5M revenue on a $9.5M Series A-sized raise; hiring is GTM-heavy.",
        disqualifiers=[], growth_flag=None,
    )
    out = ROOT / "out" / "smoke_test.xlsx"
    out.parent.mkdir(exist_ok=True)
    excel.write([saas], cfg, str(out))
    assert out.exists() and out.stat().st_size > 5000

    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb.sheetnames == ["Pipeline", "Methodology"]
    ws = wb["Pipeline"]
    assert ws["A2"].value == "Ledgerline Software Inc."
    assert ws["C2"].value == "$1M–$5M (disclosed)"
    assert ws["D2"].value == "high"

    print("smoke test: all assertions passed")


if __name__ == "__main__":
    main()
